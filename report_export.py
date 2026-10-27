from io import BytesIO
import pandas as pd
from openpyxl.utils import get_column_letter

def build_excel_report(result: pd.DataFrame, summary: pd.DataFrame) -> bytes:
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Tong quan", index=False)
        result.to_excel(writer, sheet_name="Ket qua doi soat", index=False)
        exceptions = result[result["status"] != "Khớp"]
        exceptions.to_excel(writer, sheet_name="Can kiem tra", index=False)
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column_index, column_cells in enumerate(
                sheet.iter_cols(1, sheet.max_column), start=1
            ):
                max_length = max(len(str(cell.value or "")) for cell in column_cells)
                column_letter = get_column_letter(column_index)
                sheet.column_dimensions[column_letter].width = min(max_length + 2, 32)
    return buffer.getvalue()
